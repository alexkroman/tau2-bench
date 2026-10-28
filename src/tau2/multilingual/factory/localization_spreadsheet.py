"""Map-backed language-tab workbooks and source-bound reviewer decisions.

The optional workbook dependency is loaded only by file operations. Review edits
produce feedback and exact-map approvals, never unvalidated entity replacements.
"""

import argparse
import json
import os
import tempfile
import unicodedata
from pathlib import Path
from typing import Literal

from loguru import logger
from pydantic import Field, model_validator

from tau2.multilingual.factory.localization_artifact import RegionalMapArtifact
from tau2.multilingual.localization import LocalizationModel, json_content_digest
from tau2.utils.artifacts import ReviewApproval, content_digest, write_artifact

_EDITABLE = frozenset({"decision", "proposed_entity", "reason"})


class EntityReviewRow(LocalizationModel):
    """One stable authored entity and its independently editable review fields."""

    original_entity: str = Field(
        title="Original entity", description="Canonical value."
    )
    suggested_entity: str = Field(
        title="Suggested localized entity", description="Native candidate."
    )
    decision: Literal["", "Approved", "Disapproved"] = Field(
        default="", title="Approved / Disapproved", description="Blank is pending."
    )
    proposed_entity: str = Field(
        default="",
        title="Proposed entity",
        description="Reviewer replacement suggestion, never automatically applied.",
    )
    reason: str = Field(
        default="",
        title="Reason",
        description="Explanation for disapproval or replacement.",
    )
    domain: str = Field(title="Domain", description="Owning domain.")
    category: str = Field(title="Entity type", description="Map category.")
    origin: Literal["Reused", "New"] = Field(
        title="Candidate origin", description="Recorded candidate lineage."
    )
    romanized: str = Field(
        title="Romanized reference",
        description="Reference form, not an automatically accepted answer.",
    )
    explanation: str = Field(
        title="English explanation", description="Meaning of the localized value."
    )
    context: str = Field(
        title="Context and rationale",
        description="Authored notes about use and constraints.",
    )
    lineage: str = Field(
        title="Reuse evidence",
        description="Prior artifact and approval, when available.",
    )
    language: str = Field(title="Language ID", description="Exact pack language.")
    entry_id: str = Field(title="Entity ID", description="Stable map entry ID.")
    entry_sha256: str = Field(
        title="Entry digest", description="Digest of the exact authored entry."
    )
    source_sha256: str = Field(
        title="Source digest", description="Digest of its canonical source manifest."
    )

    @model_validator(mode="after")
    def valid_decision(self):
        """Reject ambiguous approval and unexplained negative feedback."""
        if (
            self.decision == "Disapproved" or self.proposed_entity
        ) and not self.reason.strip():
            raise ValueError("Disapproval or proposed replacement requires a reason")
        if self.decision == "Approved" and self.proposed_entity:
            raise ValueError("Approved entity cannot also have a proposed replacement")
        return self

    @property
    def key(self) -> tuple[str, str, str]:
        """Stable identity independent of row position or tab ordering."""
        return self.language, self.domain, self.entry_id

    def without_decision(self) -> "EntityReviewRow":
        """Return the immutable displayed candidate and provenance."""
        return self.model_copy(update={field: "" for field in _EDITABLE})


class SpreadsheetManifest(LocalizationModel):
    """Pinned sidecar retained beside the editable workbook."""

    version: Literal["localization-spreadsheet-v1"] = "localization-spreadsheet-v1"
    map_sha256: dict[str, str] = Field(
        description="Language/domain to exact compact-map digest."
    )
    rows: list[EntityReviewRow] = Field(
        description="Original immutable review rows with blank decisions."
    )


class MapReviewApproval(ReviewApproval):
    """A completely approved map, ready for source-checked publication."""

    language: str = Field(description="Language of the approved map.")
    domain: str = Field(description="Domain of the approved map.")


class SpreadsheetReview(LocalizationModel):
    """Typed decisions and revision feedback; changes are not silently accepted."""

    version: Literal["localization-spreadsheet-review-v1"] = (
        "localization-spreadsheet-review-v1"
    )
    manifest_sha256: str = Field(description="Exact workbook source manifest reviewed.")
    reviewed_by: str = Field(min_length=1, description="Explicit reviewer attribution.")
    rows: list[EntityReviewRow] = Field(
        description="All decisions, including pending and disapproved entries."
    )
    approvals: list[MapReviewApproval] = Field(
        description="Only maps with every authored row approved unchanged."
    )


def _display(value) -> str:
    return (
        value
        if isinstance(value, str)
        else json.dumps(value, ensure_ascii=False, sort_keys=True)
    )


def _manifest(maps: list[RegionalMapArtifact]) -> SpreadsheetManifest:
    if not maps:
        raise ValueError("Review requires at least one map")
    digests = {}
    rows = []
    for mapping in sorted(maps, key=lambda m: (m.settings.language, m.source.domain)):
        key = f"{mapping.settings.language}/{mapping.source.domain}"
        if key in digests:
            raise ValueError(f"Duplicate language/domain map: {key}")
        if not mapping.entries:
            raise ValueError(f"Review requires authored entries: {key}")
        digests[key] = content_digest(mapping)
        for entry in sorted(mapping.entries, key=lambda entry: entry.id):
            rows.append(
                EntityReviewRow(
                    original_entity=_display(entry.source),
                    suggested_entity=_display(entry.forms.native),
                    domain=mapping.source.domain,
                    category=entry.category,
                    origin="Reused" if entry.reuse else "New",
                    romanized=_display(entry.forms.romanized),
                    explanation=entry.forms.english,
                    context=entry.notes,
                    lineage=entry.reuse.model_dump_json() if entry.reuse else "",
                    language=mapping.settings.language,
                    entry_id=entry.id,
                    entry_sha256=content_digest(entry),
                    source_sha256=content_digest(mapping.source),
                )
            )
        for binding in sorted(
            mapping.prose_bindings, key=lambda b: (b.entry_id, b.source)
        ):
            identifier = "prose-reference:" + json_content_digest(
                {"entry_id": binding.entry_id, "source": binding.source}
            )
            rows.append(
                EntityReviewRow(
                    original_entity=binding.source,
                    suggested_entity=_display(binding.forms.native),
                    domain=mapping.source.domain,
                    category="prose_reference",
                    origin="Reused" if binding.kind == "reuse" else "New",
                    romanized=_display(binding.forms.romanized),
                    explanation=binding.forms.english,
                    context=f"{binding.kind} reference in {binding.entry_id}: {binding.reason}",
                    lineage=binding.model_dump_json(
                        exclude={"source", "forms", "reason"}
                    ),
                    language=mapping.settings.language,
                    entry_id=identifier,
                    entry_sha256=content_digest(binding),
                    source_sha256=content_digest(mapping.source),
                )
            )
    if len({row.key for row in rows}) != len(rows):
        raise ValueError("Duplicate entity IDs in map")
    return SpreadsheetManifest(map_sha256=digests, rows=rows)


def _sidecar(path: Path) -> Path:
    return path.with_suffix(".manifest.json")


def _row_height(values: list[str], widths: list[int]) -> float:
    """Size human review text; machine lineage must not inflate every reused row."""
    lines = 1
    for value, width in zip(values[:11], widths[:11], strict=True):
        usable = width - 2
        wrapped = 0
        for line in value.splitlines() or [""]:
            length = sum(
                2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1
                for char in line
            )
            wrapped += max(1, (length + usable - 1) // usable)
        lines = max(lines, wrapped)
    return min(409.5, max(72, (lines + 1) * 15))


def _openpyxl():
    try:
        import openpyxl
    except ImportError as error:
        raise ImportError(
            "Spreadsheet review requires uv sync --extra localization-review"
        ) from error
    return openpyxl


def _read_rows(path: Path, manifest: SpreadsheetManifest) -> list[EntityReviewRow]:
    expected = {row.key: row for row in manifest.rows}
    columns = EntityReviewRow.model_fields
    headers = [field.title for field in columns.values()]
    workbook = _openpyxl().load_workbook(path, data_only=False)
    rows = {}
    try:
        if set(workbook.sheetnames) != {row.language for row in manifest.rows}:
            raise ValueError("Review language tabs differ from manifest")
        for sheet in workbook:
            if [
                sheet.cell(1, col).value for col in range(1, len(headers) + 1)
            ] != headers:
                raise ValueError("Review headers differ from model contract")
            if sheet.max_column > len(headers):
                raise ValueError("Unexpected review columns")
            for cells in sheet.iter_rows(min_row=2, max_col=len(headers)):
                if all(cell.value is None for cell in cells):
                    continue
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("Review cells cannot contain a formula")
                row = EntityReviewRow.model_validate(
                    {
                        field: "" if cell.value is None else cell.value
                        for field, cell in zip(columns, cells, strict=True)
                    }
                )
                if row.key in rows:
                    raise ValueError(f"Review has duplicate row: {row.key}")
                if row.language != sheet.title or row.key not in expected:
                    raise ValueError(f"Unknown review row or language: {row.key}")
                if row.without_decision() != expected[row.key]:
                    raise ValueError(
                        f"Review changed immutable candidate fields: {row.key}"
                    )
                rows[row.key] = row
        if rows.keys() != expected.keys():
            raise ValueError("Review has missing entity rows")
        return [rows[row.key] for row in manifest.rows]
    finally:
        workbook.close()


def export_localization_spreadsheet(
    maps: list[RegionalMapArtifact], output: Path, *, previous: Path | None = None
) -> SpreadsheetManifest:
    """Write one language tab per locale and retain unchanged prior decisions."""
    manifest = _manifest(maps)
    if output.suffix != ".xlsx":
        raise ValueError("Spreadsheet output must end in .xlsx")
    if output.exists():
        saved = SpreadsheetManifest.model_validate_json(_sidecar(output).read_text())
        if saved != manifest:
            raise FileExistsError("Review inputs changed; choose a new workbook path")
        _read_rows(output, manifest)
        return manifest
    decisions = {}
    if previous is not None:
        old = SpreadsheetManifest.model_validate_json(_sidecar(previous).read_text())
        if not old.map_sha256.keys() <= manifest.map_sha256.keys():
            raise ValueError("Refreshing review cannot remove existing domain maps")
        if not {row.key for row in old.rows} <= {row.key for row in manifest.rows}:
            raise ValueError("Refreshing review cannot remove existing entity rows")
        decisions = {row.key: row for row in _read_rows(previous, old)}
    openpyxl = _openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    columns = EntityReviewRow.model_fields
    widths = [34, 34, 24, 34, 48, 22, 22, 18, 30, 48, 60, 50, 14, 20, 20, 20]
    for language in sorted({row.language for row in manifest.rows}):
        sheet = workbook.create_sheet(language)
        sheet.freeze_panes = "C2"
        sheet.sheet_view.showGridLines = False
        sheet.append([field.title for field in columns.values()])
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(color="FFFFFF", bold=True)
        for row in (row for row in manifest.rows if row.language == language):
            prior = decisions.get(row.key)
            if prior is not None and prior.without_decision() == row:
                row = prior
            values = list(row.model_dump().values())
            if any(len(value) > 32767 for value in values):
                raise ValueError(
                    f"Entity text exceeds spreadsheet cell limit: {row.key}"
                )
            sheet.append(values)
            for field, cell in zip(columns, sheet[sheet.max_row], strict=True):
                cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.protection = Protection(locked=field not in _EDITABLE)
                if field in _EDITABLE:
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
            sheet.row_dimensions[sheet.max_row].height = _row_height(values, widths)
        for col, width in enumerate(widths, start=1):
            dim = sheet.column_dimensions[openpyxl.utils.get_column_letter(col)]
            dim.width = width
            dim.hidden = col >= 13
        sheet.auto_filter.ref = sheet.dimensions
        validation = DataValidation(
            type="list", formula1='"Approved,Disapproved"', allow_blank=True
        )
        validation.errorTitle = "Select a review decision"
        validation.error = "Choose Approved or Disapproved, or leave blank for pending."
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(f"C2:C{sheet.max_row}")
        sheet.protection.sheet = True
        sheet.protection.autoFilter = False
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        suffix=".xlsx", dir=output.parent, delete=False
    ) as handle:
        temp = Path(handle.name)
    try:
        workbook.save(temp)
        write_artifact(manifest, _sidecar(output))
        os.link(temp, output)
    finally:
        workbook.close()
        temp.unlink(missing_ok=True)
    return manifest


def ingest_localization_spreadsheet(
    maps: list[RegionalMapArtifact], workbook: Path, reviewed_by: str
) -> SpreadsheetReview:
    """Validate exact source/candidate coverage and return typed review evidence."""
    manifest = _manifest(maps)
    saved = SpreadsheetManifest.model_validate_json(_sidecar(workbook).read_text())
    if manifest != saved:
        raise ValueError("Review is stale: map inputs differ from exported manifest")
    if not reviewed_by.strip():
        raise ValueError("Reviewer identity is required")
    rows = _read_rows(workbook, manifest)
    approvals = []
    for mapping in maps:
        selected = [
            row
            for row in rows
            if row.language == mapping.settings.language
            and row.domain == mapping.source.domain
        ]
        if selected and all(row.decision == "Approved" for row in selected):
            approvals.append(
                MapReviewApproval(
                    language=mapping.settings.language,
                    domain=mapping.source.domain,
                    content_sha256=content_digest(mapping),
                    reviewed_by=reviewed_by,
                )
            )
    return SpreadsheetReview(
        manifest_sha256=content_digest(manifest),
        reviewed_by=reviewed_by,
        rows=rows,
        approvals=approvals,
    )


def _source_checked_maps(paths: list[Path]) -> list[RegionalMapArtifact]:
    from tau2.multilingual.factory.localization_artifact import read_localization_map

    result = []
    for path in paths:
        read_localization_map(path)
        result.append(RegionalMapArtifact.model_validate_json(path.read_text()))
    return result


def run_spreadsheet_export(args: argparse.Namespace) -> None:
    """Export validated compact maps, without generating or approving content."""
    result = export_localization_spreadsheet(
        _source_checked_maps(args.map), args.output, previous=args.previous
    )
    logger.info("Wrote {} review rows to {}", len(result.rows), args.output)


def run_spreadsheet_ingest(args: argparse.Namespace) -> None:
    """Persist decisions and feedback; only complete maps receive approvals."""
    result = ingest_localization_spreadsheet(
        _source_checked_maps(args.map), args.workbook, args.reviewed_by
    )
    write_artifact(result, args.output)
    if args.approval_dir is not None:
        for approval in result.approvals:
            write_artifact(
                ReviewApproval(
                    content_sha256=approval.content_sha256,
                    reviewed_by=approval.reviewed_by,
                ),
                args.approval_dir
                / approval.language
                / f"{approval.domain}.approval.json",
            )
    logger.info(
        "Ingested {} rows; {} maps approved", len(result.rows), len(result.approvals)
    )


def add_localization_spreadsheet_args(commands) -> None:
    """Register workbook stages under the existing factory command group."""
    export = commands.add_parser(
        "export-localization-spreadsheet",
        help="Export map-backed entity review with one tab per language",
    )
    export.add_argument("--map", type=Path, action="append", required=True)
    export.add_argument("--output", type=Path, required=True)
    export.add_argument("--previous", type=Path)
    export.set_defaults(func=run_spreadsheet_export)
    ingest = commands.add_parser(
        "ingest-localization-spreadsheet",
        help="Validate workbook decisions and export map review evidence",
    )
    ingest.add_argument("--map", type=Path, action="append", required=True)
    ingest.add_argument("--workbook", type=Path, required=True)
    ingest.add_argument("--reviewed-by", required=True)
    ingest.add_argument("--output", type=Path, required=True)
    ingest.add_argument("--approval-dir", type=Path)
    ingest.set_defaults(func=run_spreadsheet_ingest)
